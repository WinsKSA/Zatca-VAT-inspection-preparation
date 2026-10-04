"""Create the Arabic edition of the ZATCA VAT Inspection Pack.

Reads ZATCA_VAT_Inspection_Pack.xlsx (built by build_zatca_pack.py) and writes
ZATCA_VAT_Inspection_Pack_AR.xlsx: Arabic sheet names, right-to-left sheets and
Arabic labels, dropdown values, statuses and warnings. Formula logic is unchanged.
"""
import os
import re
import sys
from openpyxl import load_workbook
from openpyxl.comments import Comment

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "ZATCA_VAT_Inspection_Pack.xlsx")
OUT = sys.argv[2] if len(sys.argv) > 2 else os.path.join(HERE, "ZATCA_VAT_Inspection_Pack_AR.xlsx")

SHEETS = {
    "Dashboard": "لوحة_المتابعة", "Setup": "الإعدادات", "Checklist": "قائمة_المستندات",
    "TB_Monthly": "ميزان_المراجعة", "VAT_Returns": "الإقرارات", "Sales_Register": "سجل_المبيعات",
    "Purchase_Register": "سجل_المشتريات", "ZATCA_Payments": "المدفوعات", "Recon_Sales": "مطابقة_المبيعات",
    "Recon_OutputVAT": "مطابقة_المخرجات", "Recon_InputVAT": "مطابقة_المدخلات",
    "Recon_VAT_Account": "مطابقة_حساب_الضريبة", "Inspection_Guide": "دليل_الفحص", "Lists": "القوائم",
}
S = SHEETS  # short alias for use inside the strings below

# Dropdown values: identical to the browser app so the app can import this workbook.
LIST_VALUES = {
    "Standard 15% (Box 1)": "النسبة الأساسية 15% (خانة 1)",
    "Citizens - Healthcare/Education (Box 2)": "المواطنون - صحة/تعليم (خانة 2)",
    "Zero-rated domestic (Box 3)": "محلية بنسبة صفر (خانة 3)",
    "Export (Box 4)": "صادرات (خانة 4)",
    "Exempt (Box 5)": "معفاة (خانة 5)",
    "Out of scope": "خارج النطاق",
    "Standard domestic 15% (Box 7)": "محلية بالنسبة الأساسية 15% (خانة 7)",
    "Import - VAT paid at customs (Box 8)": "استيراد - ضريبة مدفوعة في الجمارك (خانة 8)",
    "Import - Reverse charge (Box 9)": "استيراد - احتساب عكسي (خانة 9)",
    "Zero-rated purchase (Box 10)": "مشتريات بنسبة صفر (خانة 10)",
    "Exempt purchase (Box 11)": "مشتريات معفاة (خانة 11)",
    "Out of scope / Not reported": "خارج النطاق / غير مُقر عنها",
    "Revenue - Out of scope / Other income": "إيرادات خارج النطاق / إيرادات أخرى",
    "Output VAT": "ضريبة المخرجات", "Input VAT": "ضريبة المدخلات",
    "VAT Payable / Settlement": "ضريبة مستحقة / تسوية", "Not VAT relevant": "غير متعلق بالضريبة",
    "Tax Invoice": "فاتورة ضريبية", "Simplified Tax Invoice": "فاتورة ضريبية مبسطة",
    "Credit Note": "إشعار دائن", "Debit Note": "إشعار مدين",
    "B2B - VAT registered": "منشأة - مسجلة في الضريبة", "B2B - Not registered": "منشأة - غير مسجلة",
    "B2C - Individual": "فرد (مستهلك)", "Government": "جهة حكومية",
    "General business expense": "مصروف تشغيلي عام", "Capital asset": "أصل رأسمالي",
    "Entertainment / hospitality (blocked)": "ترفيه / ضيافة (غير قابلة للخصم)",
    "Motor vehicle - private use (blocked)": "سيارة - استخدام شخصي (غير قابلة للخصم)",
    "Staff personal benefit (blocked)": "منفعة شخصية للموظفين (غير قابلة للخصم)",
    "Other non-business (blocked)": "أخرى غير تجارية (غير قابلة للخصم)",
    "Yes": "نعم", "No": "لا",
    "Not started": "لم يبدأ", "In progress": "قيد التنفيذ", "Ready": "جاهز", "Submitted": "تم التسليم", "N/A": "لا ينطبق",
    "Asset": "أصول", "Liability": "التزامات", "Equity": "حقوق ملكية", "Revenue": "إيرادات", "Expense": "مصروفات",
    "Monthly": "شهري", "Quarterly": "ربع سنوي",
}

T = dict(LIST_VALUES)
T.update({
    # literals used inside formulas
    "blocked": "غير قابلة للخصم", "Saudi Arabia": "المملكة العربية السعودية", "Q": "الربع ", "Prior period": "فترة سابقة",
    "Company: ": "المنشأة: ", "   |   VAT No: ": "   |   الرقم الضريبي: ", "   |   Period: ": "   |   الفترة: ",
    " to ": " إلى ", "   |   Filing: ": "   |   الإقرار: ", " days": " يوم",
    "Outside audit periods; ": "التاريخ خارج فترات الفحص؛ ", "No VAT category; ": "لا توجد فئة ضريبية؛ ",
    "VAT ≠ 15% × net; ": "الضريبة ≠ 15% × الصافي؛ ", "Invalid/missing customer VAT no; ": "الرقم الضريبي للعميل غير صحيح أو مفقود؛ ",
    "Not reported/cleared on FATOORA; ": "غير مُبلّغ عنها في فاتورة؛ ", "Zero-rate/export evidence missing; ": "إثبات النسبة الصفرية أو التصدير مفقود؛ ",
    "Credit note should be negative; ": "يجب أن يكون الإشعار الدائن بالسالب؛ ",
    "Simplified invoice issued to VAT-registered B2B; ": "فاتورة مبسطة لمنشأة مسجلة في الضريبة؛ ",
    "Invalid/missing supplier VAT no; ": "الرقم الضريبي للمورد غير صحيح أو مفقود؛ ",
    "VAT claimed without valid tax invoice; ": "ضريبة مخصومة دون فاتورة ضريبية صحيحة؛ ",
    "Blocked input VAT claimed; ": "خصم ضريبة غير قابلة للخصم؛ ", "Customs declaration no missing; ": "رقم البيان الجمركي مفقود؛ ",
    "RCM used for a local supplier; ": "احتساب عكسي على مورد محلي؛ ", "VAT claimed on non-taxable category; ": "خصم ضريبة على فئة غير خاضعة؛ ",
    "– No activity": "– لا يوجد نشاط", "– No filing date": "– لا يوجد تاريخ تقديم", "– Not entered": "– لم يُدخل",
    "✓ Agrees": "✓ مطابق", "✓ All mapped": "✓ جميع الحسابات مصنفة", "✓ Balanced": "✓ متوازن", "✓ Clean": "✓ سليم",
    "✓ Fully reconciled": "✓ مطابق بالكامل", "✓ Low risk": "✓ مخاطر منخفضة", "✓ Matched": "✓ مطابق", "✓ None": "✓ لا يوجد",
    "✓ OK": "✓ سليم", "✓ On time": "✓ في الموعد", "✓ Ready": "✓ جاهز", "✓ Reconciled": "✓ مطابق", "✓ Settled": "✓ مسدد",
    "✓ Valid format": "✓ صيغة صحيحة", "✗ >5,000 – needs voluntary disclosure": "✗ يتجاوز 5,000 – يتطلب إفصاحاً طوعياً",
    "✗ Check": "✗ يحتاج مراجعة", "✗ Explain": "✗ يحتاج تفسيراً", "✗ Exposure": "✗ يوجد تعرض", "✗ Filed ≠ recalc": "✗ المقدم ≠ المعاد احتسابه",
    "✗ Fix TB": "✗ صحح الميزان", "✗ Incomplete": "✗ غير مكتمل", "✗ Invalid VAT number format": "✗ صيغة الرقم الضريبي غير صحيحة",
    "✗ Investigate": "✗ يحتاج إلى فحص", "✗ Late": "✗ متأخر", "✗ Late ": "✗ متأخر ", "✗ Map accounts": "✗ صنّف الحسابات",
    "✗ Prepare defence / consider voluntary disclosure": "✗ جهّز الدفاع / ادرس الإفصاح الطوعي",
    "✗ Review Recon_Sales": f"✗ راجع {S['Recon_Sales']}", "✗ Review flags": "✗ راجع الملاحظات",
    "✗ Unexplained – ZATCA will treat as undeclared sales": "✗ غير مفسَّر – ستعتبره الهيئة مبيعات غير مُقر بها", "✗ Unpaid": "✗ غير مسدد",
    # Dashboard
    "ZATCA VAT Inspection Pack – Dashboard": "حزمة الفحص الضريبي – لوحة المتابعة", "Indicator": "المؤشر", "Value (SAR)": "القيمة (ريال)",
    "Status": "الحالة", "What it means": "المعنى", "1. Sales & Revenue": "1. المبيعات والإيرادات", "2. Output VAT": "2. ضريبة المخرجات",
    "3. Input VAT": "3. ضريبة المدخلات", "4. VAT Account, Filing & Payment": "4. حساب الضريبة والتقديم والسداد", "5. Data Quality": "5. جودة البيانات",
    "6. Indicative Exposure (before objection / penalty relief)": "6. التعرض التقديري (قبل الاعتراض أو الإعفاء من الغرامات)",
    "Total sales per VAT returns (Box 6)": "إجمالي المبيعات حسب الإقرارات (خانة 6)", "Revenue subject to VAT per Trial Balance": "الإيرادات الخاضعة للضريبة حسب الميزان",
    "Unexplained revenue difference": "فرق الإيرادات غير المفسَّر", "ZATCA treats unexplained excess revenue as undeclared taxable sales": "تعتبر الهيئة الزيادة غير المفسرة في الإيرادات مبيعات خاضعة غير مُقر بها",
    "Sales periods/boxes not matched": "فترات/خانات مبيعات غير مطابقة", "Output VAT per returns": "ضريبة المخرجات حسب الإقرارات",
    "Output VAT per GL": "ضريبة المخرجات حسب الدفاتر", "Potential under-declared output VAT": "ضريبة مخرجات محتمل عدم الإقرار بها",
    "Higher of register or GL above the return, per period": "الأعلى من السجل أو الدفاتر فوق الإقرار، لكل فترة",
    "Input VAT claimed per returns": "ضريبة المدخلات المخصومة حسب الإقرارات", "Defensible input VAT per register": "ضريبة المدخلات القابلة للدفاع حسب السجل",
    "Potential over-claimed input VAT": "ضريبة مدخلات محتمل خصمها بالزيادة", "Claimed in return but not supported by a valid, non-blocked invoice": "مخصومة في الإقرار دون فاتورة صحيحة قابلة للخصم",
    "Returns filed late": "إقرارات مقدمة متأخرة", "Return arithmetic errors": "أخطاء حسابية في الإقرارات", "VAT unpaid (outstanding)": "ضريبة غير مسددة",
    "GL VAT account periods not agreeing": "فترات حساب الضريبة غير المطابقة", "Est. late payment penalties": "غرامات التأخر في السداد التقديرية",
    "5% per month or part on late / unpaid tax": "5% عن كل شهر أو جزء منه على الضريبة المتأخرة أو غير المسددة",
    "Sales rows with exceptions": "صفوف مبيعات بها ملاحظات", "Purchase rows with exceptions": "صفوف مشتريات بها ملاحظات",
    "TB months out of balance": "أشهر غير متوازنة في الميزان", "TB accounts without VAT mapping": "حسابات بدون تصنيف ضريبي",
    "Document checklist readiness": "جاهزية قائمة المستندات",
    "Tax exposure (under-declared + over-claimed + unexplained revenue × rate)": "التعرض الضريبي (نقص الإقرار + الخصم بالزيادة + الإيرادات غير المفسرة × النسبة)",
    "Unexplained revenue (TB > returns) × 15%. Indicative only": "الإيرادات غير المفسرة (الميزان > الإقرارات) × 15%. تقديري فقط",
    "Incorrect-return penalty on tax exposure": "غرامة الإقرار الخاطئ على التعرض الضريبي", "Art. 42 – 50% of the difference (indicative)": "المادة 42 – 50% من الفرق (تقديري)",
    "Late payment penalties (est.)": "غرامات التأخر في السداد (تقديرية)", "TOTAL INDICATIVE EXPOSURE": "إجمالي التعرض التقديري",
    "Legend: yellow cells with blue text = inputs · grey = formulas · green text = links to other sheets · ✓ = OK · ✗ = needs action. All amounts in SAR. Example rows (Jan 2025) are included to show the format – overwrite them with your data.":
        "دليل الألوان: الخلايا الصفراء بخط أزرق = مدخلات · الرمادية = معادلات · الخط الأخضر = روابط لأوراق أخرى · ✓ = سليم · ✗ = يحتاج إجراء. جميع المبالغ بالريال. صفوف المثال (يناير 2025) لتوضيح الصيغة فقط – استبدلها ببياناتك.",
    # Setup
    "Setup – Company & Audit Period": "الإعدادات – المنشأة وفترة الفحص",
    "Yellow cells with blue text are inputs. Everything else in the workbook is driven by these settings.": "الخلايا الصفراء بخط أزرق هي المدخلات. باقي الملف يعتمد على هذه الإعدادات.",
    "Setting": "الإعداد", "Value": "القيمة", "Check / Note": "الفحص / ملاحظة", "Company Name (اسم المنشأة)": "اسم المنشأة",
    "Example Trading Company": "شركة المثال للتجارة", "Your legal name as on the VAT certificate": "الاسم النظامي كما في شهادة التسجيل الضريبي",
    "VAT Registration No (الرقم الضريبي)": "الرقم الضريبي", "Commercial Registration No": "رقم السجل التجاري",
    "Fiscal Year Start (audit year)": "بداية السنة المالية (سنة الفحص)", "First day of the year under inspection": "أول يوم في السنة محل الفحص",
    "Filing Frequency": "دورية تقديم الإقرار", "Monthly if taxable supplies > SAR 40M/yr, otherwise Quarterly": "شهري إذا تجاوزت التوريدات الخاضعة 40 مليون ريال سنوياً، وإلا ربع سنوي",
    "Standard VAT Rate": "نسبة الضريبة الأساسية", "15% since 1 July 2020": "15% منذ 1 يوليو 2020", "Variance Tolerance (SAR)": "حد فروق التقريب (ريال)",
    "Differences within this amount are treated as rounding": "الفروق ضمن هذا المبلغ تعتبر فروق تقريب", "As-of Date (for penalty estimates)": "تاريخ احتساب الغرامات",
    "Usually today's date or the inspection date": "عادةً تاريخ اليوم أو تاريخ الفحص", "Late Payment Penalty (per month or part)": "غرامة التأخر في السداد (عن كل شهر أو جزء منه)",
    "VAT Law Art. 43 – verify current rule": "المادة 43 من النظام – تحقق من الحكم الحالي", "Incorrect Return Penalty (% of difference)": "غرامة الإقرار الخاطئ (% من الفرق)",
    "VAT Law Art. 42 – indicative, verify current rule": "المادة 42 من النظام – تقديري، تحقق من الحكم الحالي",
    "Periods in Year (calc)": "عدد الفترات في السنة (محسوب)", "Months per Period (calc)": "عدد الأشهر في الفترة (محسوب)",
    "Tax Periods (auto-generated from the settings above)": "الفترات الضريبية (تُنشأ تلقائياً من الإعدادات أعلاه)",
    "Period #": "رقم الفترة", "Period Label": "الفترة", "Start Date": "تاريخ البداية", "End Date": "تاريخ النهاية", "Return & Payment Due Date": "موعد الإقرار والسداد",
    # Checklist
    "Document Request Checklist (قائمة المستندات المطلوبة)": "قائمة المستندات المطلوبة",
    "Typical documents ZATCA requests in a VAT inspection. Track status, owner and file reference for each.": "المستندات التي تطلبها الهيئة عادةً في فحص ضريبة القيمة المضافة. تابع الحالة والمسؤول ومرجع الملف لكل بند.",
    "Area": "المجال", "Document / Analysis Requested": "المستند / التحليل المطلوب", "Prepared in this workbook?": "مُعدّ في هذا الملف؟", "Owner": "المسؤول",
    "Due Date": "تاريخ الاستحقاق", "File Ref / Location": "مرجع الملف / مكانه", "Notes": "ملاحظات", "Readiness:": "الجاهزية:",
    "General": "عام", "Financial": "مالي", "VAT Returns": "الإقرارات الضريبية", "Reconciliation": "مطابقة", "Sales": "المبيعات", "Purchases": "المشتريات",
    "Cross-tax": "ضرائب أخرى", "E-invoicing": "الفوترة الإلكترونية", "Response": "الرد على الهيئة",
    "VAT registration certificate, Commercial Registration, Articles of Association": "شهادة التسجيل في ضريبة القيمة المضافة، السجل التجاري، عقد التأسيس",
    "Description of business activities and revenue streams + VAT treatment memo for each": "وصف الأنشطة ومصادر الإيرادات مع مذكرة المعالجة الضريبية لكل منها",
    "List of branches, related parties and VAT group members (if any)": "قائمة الفروع والأطراف ذات العلاقة وأعضاء المجموعة الضريبية (إن وجدت)",
    "Chart of accounts with VAT mapping": "دليل الحسابات مع التصنيف الضريبي",
    "ERP / accounting system description and e-invoicing (FATOORA) solution details, onboarding & CSID": "وصف نظام ERP / المحاسبة وتفاصيل حل الفوترة الإلكترونية (فاتورة) والربط وشهادة التشفير",
    "Authorisation letter for the representative dealing with ZATCA": "خطاب تفويض الممثل المسؤول عن التعامل مع الهيئة",
    "Audited financial statements for each year under review": "القوائم المالية المدققة لكل سنة محل الفحص",
    "Monthly trial balances (opening, movements, closing)": "موازين المراجعة الشهرية (افتتاحي، حركات، ختامي)",
    "General ledger detail: revenue, VAT, and major expense accounts": "تفاصيل دفتر الأستاذ: الإيرادات والضريبة والمصروفات الرئيسية",
    "Bank statements (ZATCA compares receipts with declared sales)": "كشوف الحسابات البنكية (تقارن الهيئة المتحصلات بالمبيعات المُقر بها)",
    "Copies of all filed VAT returns and SADAD payment receipts": "نسخ جميع الإقرارات المقدمة وإيصالات السداد",
    "Voluntary disclosures / corrections filed (Box 14 usage)": "الإفصاحات الطوعية والتصحيحات المقدمة (استخدام خانة 14)",
    "Revenue per financial statements vs total sales per VAT returns": "الإيرادات في القوائم المالية مقابل إجمالي المبيعات في الإقرارات",
    "Sales per VAT box vs sales listing vs GL – by period": "المبيعات لكل خانة مقابل سجل المبيعات ودفتر الأستاذ – لكل فترة",
    "Output VAT per GL vs returns": "ضريبة المخرجات في الدفاتر مقابل الإقرارات", "Input VAT per GL vs returns": "ضريبة المدخلات في الدفاتر مقابل الإقرارات",
    "VAT control (payable) account roll-forward": "حركة حساب الضريبة المستحقة", "Customs imports (Bayan data) vs Box 8": "بيانات الاستيراد الجمركية (البيان) مقابل خانة 8",
    "Sales listing (invoice level) with customer VAT numbers": "سجل المبيعات على مستوى الفاتورة مع الأرقام الضريبية للعملاء",
    "Sample tax invoices, simplified invoices, credit/debit notes (XML / QR code)": "عينات من الفواتير الضريبية والمبسطة والإشعارات الدائنة والمدينة (XML / رمز QR)",
    "Export evidence: customs exit declarations, bills of lading, contracts, proof of payment": "إثباتات التصدير: بيانات الخروج الجمركية، بوالص الشحن، العقود، إثبات التحصيل",
    "Zero-rated domestic supply evidence (e.g. qualifying medicines, intl. transport)": "إثباتات التوريدات المحلية بنسبة صفر (مثل الأدوية المؤهلة والنقل الدولي)",
    "Exempt supplies support (financial services margin, residential leases)": "مستندات التوريدات المعفاة (هامش الخدمات المالية، الإيجارات السكنية)",
    "Major customer contracts incl. government contracts and VAT clauses": "العقود الرئيسية مع العملاء بما فيها العقود الحكومية وبنود الضريبة",
    "Deemed supplies: free samples, gifts, own use, staff benefits, write-offs": "التوريدات الافتراضية: العينات المجانية، الهدايا، الاستخدام الخاص، منافع الموظفين، الشطب",
    "Related-party / intercompany charges and their VAT treatment": "تعاملات الأطراف ذات العلاقة والمعالجة الضريبية لها",
    "Fixed asset & real estate disposals (VAT / RETT treatment)": "استبعاد الأصول الثابتة والعقارات (معالجة ضريبة القيمة المضافة / ضريبة التصرفات العقارية)",
    "Purchase listing (invoice level) with supplier VAT numbers": "سجل المشتريات على مستوى الفاتورة مع الأرقام الضريبية للموردين",
    "Sample supplier tax invoices supporting input VAT claimed": "عينات من فواتير الموردين المؤيدة لضريبة المدخلات المخصومة",
    "Import customs declarations (Bayan) and customs VAT payment proof": "البيانات الجمركية للاستيراد وإثبات سداد الضريبة في الجمارك",
    "Reverse charge: foreign supplier invoices, contracts, RCM calculation": "الاحتساب العكسي: فواتير الموردين الأجانب والعقود وطريقة الاحتساب",
    "Blocked input VAT analysis (entertainment, vehicles, staff benefits)": "تحليل ضريبة المدخلات غير القابلة للخصم (الترفيه، السيارات، منافع الموظفين)",
    "Fixed asset register and capital assets adjustment": "سجل الأصول الثابتة وتعديلات الأصول الرأسمالية",
    "Input VAT apportionment calculation (if exempt supplies exist)": "احتساب نسبة توزيع ضريبة المدخلات (في حال وجود توريدات معفاة)",
    "Withholding tax returns (foreign payments should match RCM in Box 9)": "إقرارات ضريبة الاستقطاع (يجب أن تطابق المدفوعات للخارج خانة 9)",
    "Zakat / income tax return revenue vs VAT revenue": "إيرادات إقرار الزكاة / ضريبة الدخل مقابل إيرادات ضريبة القيمة المضافة",
    "E-invoicing compliance: Phase 2 integration wave, cleared/reported invoices report": "الالتزام بالفوترة الإلكترونية: موجة الربط في المرحلة الثانية وتقرير الفواتير المعتمدة / المبلّغة",
    "Written explanations / memos for every difference found": "تفسيرات ومذكرات مكتوبة لكل فرق يتم اكتشافه",
    f"TB_Monthly col D": f"{S['TB_Monthly']} عمود D", "TB_Monthly": S["TB_Monthly"],
    "VAT_Returns, ZATCA_Payments": f"{S['VAT_Returns']}، {S['ZATCA_Payments']}", "VAT_Returns col AJ": f"{S['VAT_Returns']} عمود AJ",
    "Recon_Sales (bottom)": f"{S['Recon_Sales']} (أسفل الورقة)", "Recon_Sales": S["Recon_Sales"], "Recon_OutputVAT": S["Recon_OutputVAT"],
    "Recon_InputVAT": S["Recon_InputVAT"], "Recon_VAT_Account": S["Recon_VAT_Account"], "Recon_InputVAT cols L–N": f"{S['Recon_InputVAT']} الأعمدة L–N",
    "Sales_Register": S["Sales_Register"], "Sales_Register col N": f"{S['Sales_Register']} عمود N", "Sales_Register col L": f"{S['Sales_Register']} عمود L",
    "Purchase_Register": S["Purchase_Register"], "Purchase_Register col K": f"{S['Purchase_Register']} عمود K", "Purchase_Register col M": f"{S['Purchase_Register']} عمود M",
    "Purchase_Register (Box 9)": f"{S['Purchase_Register']} (خانة 9)", "Purchase_Register col G": f"{S['Purchase_Register']} عمود G", "Recon sheets": "أوراق المطابقة",
    # TB
    "Monthly Trial Balance (ميزان المراجعة الشهري)": "ميزان المراجعة الشهري",
    "Paste your TB: opening balance + monthly NET movements. Debit = positive, Credit = negative. Map every account in column D.": "الصق الميزان: الرصيد الافتتاحي + صافي الحركة الشهرية. المدين بالموجب والدائن بالسالب. صنّف كل حساب في العمود D.",
    "Balance check (must = 0) →": "فحص التوازن (يجب = 0) ←", "Month-end →": "نهاية الشهر ←", "Account Code": "رمز الحساب", "Account Name": "اسم الحساب",
    "FS Class": "تصنيف القوائم المالية", "VAT Mapping": "التصنيف الضريبي", "Opening Balance": "الرصيد الافتتاحي", "Total Movements": "إجمالي الحركات", "Closing Balance": "الرصيد الختامي",
    "Cash at Bank": "النقد لدى البنوك", "VAT Payable - Settlement with ZATCA": "ضريبة مستحقة - تسوية مع الهيئة", "Share Capital / Retained Earnings": "رأس المال / الأرباح المبقاة",
    "Sales - Local": "مبيعات محلية", "Sales - Export": "مبيعات تصدير", "Other Income - Dividends": "إيرادات أخرى - توزيعات أرباح",
    "Cost of Goods Purchased": "تكلفة البضاعة المشتراة", "Consulting Fees - Foreign": "أتعاب استشارات - خارجية",
    # VAT returns
    "VAT Returns as Filed (إقرارات ضريبة القيمة المضافة)": "الإقرارات الضريبية كما قُدمت",
    "Copy each return EXACTLY as filed on the ZATCA portal. Adjustments (credit notes) as negative. Grey columns recalculate the return and check it.": "انقل كل إقرار كما قُدم في بوابة الهيئة تماماً. التعديلات (الإشعارات الدائنة) بالسالب. الأعمدة الرمادية تعيد احتساب الإقرار وتفحصه.",
    "Period": "الفترة", "Amount": "المبلغ", "Adjustment": "التعديل", "VAT": "الضريبة",
    "Box 1 – Standard rated sales 15%\nالمبيعات الخاضعة للنسبة الأساسية": "خانة 1 – المبيعات الخاضعة للنسبة الأساسية 15%",
    "Box 2 – Sales to citizens (private health / education)\nالمبيعات للمواطنين": "خانة 2 – المبيعات للمواطنين (الصحة والتعليم الأهلي)",
    "Box 3 – Zero-rated domestic sales\nالمبيعات المحلية بنسبة صفر": "خانة 3 – المبيعات المحلية الخاضعة للنسبة الصفرية",
    "Box 4 – Exports\nالصادرات": "خانة 4 – الصادرات", "Box 5 – Exempt sales\nالمبيعات المعفاة": "خانة 5 – المبيعات المعفاة",
    "Box 7 – Standard rated domestic purchases 15%\nالمشتريات المحلية الخاضعة للنسبة الأساسية": "خانة 7 – المشتريات المحلية الخاضعة للنسبة الأساسية 15%",
    "Box 8 – Imports – VAT paid at customs\nالاستيرادات – ضريبة مدفوعة في الجمارك": "خانة 8 – الاستيرادات الخاضعة للضريبة المدفوعة في الجمارك",
    "Box 9 – Imports – reverse charge (RCM)\nالاستيرادات – الاحتساب العكسي": "خانة 9 – الاستيرادات الخاضعة لآلية الاحتساب العكسي",
    "Box 10 – Zero-rated purchases\nالمشتريات بنسبة صفر": "خانة 10 – المشتريات الخاضعة للنسبة الصفرية",
    "Box 11 – Exempt purchases\nالمشتريات المعفاة": "خانة 11 – المشتريات المعفاة",
    "Box 14 – Corrections from previous periods (±5,000)": "خانة 14 – تصحيحات من فترات سابقة (±5,000)", "Box 15 – VAT credit carried forward": "خانة 15 – الرصيد الدائن المرحّل",
    "Box 16 – Net VAT due AS FILED": "خانة 16 – صافي الضريبة المستحقة كما قُدم", "Box 6 – Total sales (net, recalc)": "خانة 6 – إجمالي المبيعات (صافي، معاد احتسابه)",
    "Box 6 – Output VAT incl. RCM (recalc)": "خانة 6 – ضريبة المخرجات شاملة الاحتساب العكسي (معاد احتسابها)",
    "Box 12 – Total purchases (net, recalc)": "خانة 12 – إجمالي المشتريات (صافي، معاد احتسابه)", "Box 12 – Input VAT (recalc)": "خانة 12 – ضريبة المدخلات (معاد احتسابها)",
    "Box 13 – Net VAT for period (recalc)": "خانة 13 – صافي ضريبة الفترة (معاد احتسابه)", "Box 16 – Net VAT due (recalc)": "خانة 16 – صافي الضريبة المستحقة (معاد احتسابه)",
    "Arithmetic check (filed vs recalc)": "الفحص الحسابي (المقدم مقابل المعاد احتسابه)", "Filed on time?": "قُدم في الموعد؟", "Box 14 within ±5,000?": "خانة 14 ضمن ±5,000؟",
    "TOTAL YEAR": "إجمالي السنة", "Example: January 2025 is filled in to show the format – overwrite with your real filed figures.": "مثال: تم تعبئة يناير 2025 لتوضيح الصيغة – استبدله بأرقامك الفعلية المقدمة.",
    # Sales register
    "Sales Register – invoice level (سجل المبيعات)": "سجل المبيعات – على مستوى الفاتورة",
    "Export from ERP / FATOORA. Credit notes as NEGATIVE amounts. Grey columns calculate expected VAT and flag issues ZATCA will look for.": "صدّر البيانات من نظام ERP أو فاتورة. الإشعارات الدائنة بالسالب. الأعمدة الرمادية تحسب الضريبة المتوقعة وتظهر الملاحظات التي ستبحث عنها الهيئة.",
    "Invoice Date": "تاريخ الفاتورة", "Invoice No": "رقم الفاتورة", "Document Type": "نوع المستند", "Customer Name": "اسم العميل", "Customer VAT No": "الرقم الضريبي للعميل",
    "Customer Type": "نوع العميل", "VAT Category": "الفئة الضريبية", "Net Amount (SAR)": "المبلغ الصافي (ريال)", "VAT Amount (SAR)": "مبلغ الضريبة (ريال)", "Total (SAR)": "الإجمالي (ريال)",
    "E-Invoice UUID / Ref": "المعرّف الفريد للفاتورة الإلكترونية", "Reported/Cleared on FATOORA?": "مُبلّغة / معتمدة في فاتورة؟", "GL Revenue Account": "حساب الإيراد",
    "Evidence Ref (export / zero-rate)": "مرجع الإثبات (تصدير / نسبة صفرية)", "Expected VAT": "الضريبة المتوقعة", "VAT Difference": "فرق الضريبة", "Exceptions / Flags": "الملاحظات",
    "Rows flagged:": "صفوف بها ملاحظات:", "Rows with data:": "صفوف بها بيانات:", "ABC Trading Co": "شركة أ ب ج للتجارة", "Gulf Imports LLC (UAE)": "شركة الخليج للاستيراد (الإمارات)",
    "Customs exit decl. 123456 + B/L GLF-889": "بيان خروج جمركي 123456 + بوليصة شحن GLF-889",
    # Purchase register
    "Purchase Register – invoice level (سجل المشتريات)": "سجل المشتريات – على مستوى الفاتورة",
    "Every invoice / customs declaration on which input VAT exists. Grey columns test whether the input VAT is defensible in an inspection.": "كل فاتورة أو بيان جمركي يتضمن ضريبة مدخلات. الأعمدة الرمادية تختبر ما إذا كانت ضريبة المدخلات قابلة للدفاع عنها في الفحص.",
    "Supplier Invoice No": "رقم فاتورة المورد", "Supplier Name": "اسم المورد", "Supplier VAT No": "الرقم الضريبي للمورد", "Supplier Country": "دولة المورد",
    "Expense Nature": "طبيعة المصروف", "Valid Tax Invoice / Customs Decl. Held?": "توجد فاتورة ضريبية / بيان جمركي صحيح؟", "Input VAT Claimed in Return?": "خُصمت الضريبة في الإقرار؟",
    "Customs Declaration No (Bayan)": "رقم البيان الجمركي", "GL Account": "حساب الأستاذ", "VAT Claimed": "الضريبة المخصومة", "Defensible VAT": "الضريبة القابلة للدفاع عنها",
    "VAT at Risk": "الضريبة المعرضة للرفض", "VAT at risk (SAR):": "الضريبة المعرضة للرفض (ريال):", "Al Noor Supplies": "مؤسسة النور للتوريدات",
    "Global Consulting Ltd": "شركة جلوبال للاستشارات المحدودة", "United Kingdom": "المملكة المتحدة",
    # Payments
    "Payments to / Refunds from ZATCA (المدفوعات)": "المدفوعات للهيئة / المستردات منها",
    "One row per SADAD payment. Refunds received as NEGATIVE. Use Period # 0 for payments of periods before the audit year.": "صف لكل عملية سداد. المستردات بالسالب. استخدم رقم الفترة 0 لمدفوعات فترات ما قبل سنة الفحص.",
    "Payment Date": "تاريخ السداد", "SADAD / Payment Ref": "مرجع السداد", "Amount (SAR)": "المبلغ (ريال)", "Days Late": "أيام التأخير",
    "Months Late (or part)": "أشهر التأخير (أو جزء منها)", "Est. Late Payment Penalty": "غرامة التأخر التقديرية", "Jan 2025 VAT": "ضريبة يناير 2025",
    # Recon sales
    "Sales Reconciliation – Returns vs Register vs Trial Balance": "مطابقة المبيعات – الإقرارات مقابل السجل مقابل الميزان",
    "For each VAT box: what you DECLARED vs what you INVOICED vs what you BOOKED. ZATCA starts its audit here.": "لكل خانة: ما أقررت به مقابل ما أصدرت به فواتير مقابل ما سجلته في الدفاتر. من هنا تبدأ الهيئة الفحص.",
    "Box 1 – Standard 15% (Box 1)": "خانة 1 – النسبة الأساسية 15%", "Box 2 – Citizens - Healthcare/Education (Box 2)": "خانة 2 – المواطنون (صحة / تعليم)",
    "Box 3 – Zero-rated domestic (Box 3)": "خانة 3 – محلية بنسبة صفر", "Box 4 – Export (Box 4)": "خانة 4 – الصادرات", "Box 5 – Exempt (Box 5)": "خانة 5 – المعفاة",
    "Per VAT Return (net)": "حسب الإقرار (صافي)", "Per Sales Register": "حسب سجل المبيعات", "Per Trial Balance": "حسب ميزان المراجعة",
    "Return − Register": "الإقرار − السجل", "Return − TB": "الإقرار − الميزان", "TOTAL": "الإجمالي",
    "Annual Revenue Reconciliation – Financial Statements vs VAT Returns (مطابقة الإيرادات)": "مطابقة الإيرادات السنوية – القوائم المالية مقابل الإقرارات",
    "Total revenue per Trial Balance (all accounts with FS Class = Revenue)": "إجمالي الإيرادات حسب الميزان (جميع الحسابات المصنفة إيرادات)",
    "Less: revenue mapped as Out of scope / Other income": "يُخصم: إيرادات مصنفة خارج النطاق / إيرادات أخرى",
    "Revenue subject to VAT reporting per TB": "الإيرادات الخاضعة للإقرار حسب الميزان", "Total sales per VAT returns (Box 6, full year)": "إجمالي المبيعات حسب الإقرارات (خانة 6، السنة كاملة)",
    "Difference before reconciling items": "الفرق قبل بنود التسوية",
    "Explained reconciling items (enter, with + / − sign that reduces the difference):": "بنود التسوية المفسَّرة (أدخلها بالإشارة + / − التي تقلل الفرق):",
    "e.g. Timing – December invoices booked in January": "مثال: توقيت – فواتير ديسمبر سُجلت في يناير", "e.g. Revenue accruals not yet invoiced": "مثال: إيرادات مستحقة لم تصدر بها فواتير",
    "e.g. Advance payments invoiced (VAT due on receipt)": "مثال: دفعات مقدمة صدرت بها فواتير (الضريبة مستحقة عند الاستلام)",
    "e.g. Deemed supplies declared but not in revenue": "مثال: توريدات افتراضية مُقر بها وليست ضمن الإيرادات", "Other (describe)": "أخرى (اذكرها)",
    "UNEXPLAINED DIFFERENCE": "الفرق غير المفسَّر",
    # Recon output
    "Output VAT Reconciliation (مطابقة ضريبة المخرجات)": "مطابقة ضريبة المخرجات",
    "Output VAT declared vs invoiced vs booked in GL vs 15% recomputation. Positive 'Potential under-declaration' = likely assessment.": "ضريبة المخرجات المُقر بها مقابل الفواتير مقابل الدفاتر مقابل إعادة الاحتساب بنسبة 15%. القيمة الموجبة في «نقص محتمل في الإقرار» = تقييم مرجح.",
    "Per Return (Box 6 VAT incl. RCM)": "حسب الإقرار (ضريبة خانة 6 شاملة الاحتساب العكسي)", "Register – Sales VAT": "السجل – ضريبة المبيعات", "Register – RCM VAT": "السجل – ضريبة الاحتساب العكسي",
    "Register Total": "إجمالي السجل", "Recomputed 15% × (Box 1 + Box 9)": "إعادة الاحتساب 15% × (خانة 1 + خانة 9)", "Per GL (Output VAT accounts)": "حسب الدفاتر (حسابات ضريبة المخرجات)",
    "Return − Recomputed": "الإقرار − المعاد احتسابه", "Return − GL": "الإقرار − الدفاتر", "Potential Under-declaration": "نقص محتمل في الإقرار",
    "Note: Box 2 (citizens' private healthcare/education) VAT is borne by the State and is excluded from the 15% recomputation. GL Output VAT should include RCM output VAT; if you book RCM in a separate account, map it to 'Output VAT' too.":
        "ملاحظة: ضريبة خانة 2 (الصحة والتعليم الأهلي للمواطنين) تتحملها الدولة ولا تدخل في إعادة الاحتساب. يجب أن تشمل ضريبة المخرجات في الدفاتر ضريبة الاحتساب العكسي؛ إذا سجلتها في حساب منفصل فصنّفه أيضاً «ضريبة المخرجات».",
    # Recon input
    "Input VAT Reconciliation (مطابقة ضريبة المدخلات)": "مطابقة ضريبة المدخلات",
    "Input VAT claimed vs register vs GL, plus the portion ZATCA is likely to disallow (no valid invoice, blocked, invalid supplier VAT no).": "ضريبة المدخلات المخصومة مقابل السجل والدفاتر، والجزء المرجح أن ترفضه الهيئة (لا فاتورة صحيحة، غير قابلة للخصم، رقم ضريبي غير صحيح للمورد).",
    "Per Return (Box 12 VAT)": "حسب الإقرار (ضريبة خانة 12)", "Register – VAT Claimed": "السجل – الضريبة المخصومة", "Register – Defensible VAT": "السجل – الضريبة القابلة للدفاع عنها",
    "Register – VAT at Risk": "السجل – الضريبة المعرضة للرفض", "Per GL (Input VAT accounts)": "حسب الدفاتر (حسابات ضريبة المدخلات)", "Return − Register Claimed": "الإقرار − المخصوم حسب السجل",
    "Potential Over-claim (Return − Defensible)": "زيادة محتملة في الخصم (الإقرار − القابل للدفاع)", "Box 8 VAT per Return": "ضريبة خانة 8 حسب الإقرار",
    "Customs VAT per Register": "ضريبة الجمارك حسب السجل", "Box 8 Difference": "فرق خانة 8", "Box 9 VAT per Return": "ضريبة خانة 9 حسب الإقرار",
    "RCM VAT per Register": "ضريبة الاحتساب العكسي حسب السجل", "Box 9 Difference": "فرق خانة 9",
    "Box 8 should also be matched to ZATCA's customs import data (request the importer statement from ZATCA/customs). Customs VAT is only deductible when the company is the importer of record on the Bayan.":
        "يجب أيضاً مطابقة خانة 8 مع بيانات الاستيراد الجمركية لدى الهيئة (اطلب كشف المستورد). ضريبة الجمارك قابلة للخصم فقط إذا كانت المنشأة هي المستورد المسجل في البيان الجمركي.",
    # Recon VAT account
    "VAT Control Account & Payments (مطابقة حساب الضريبة والسداد)": "مطابقة حساب الضريبة والسداد",
    "Proves the GL VAT liability = returns filed − payments made. Also estimates late filing / payment exposure.": "يثبت أن التزام الضريبة في الدفاتر = الإقرارات المقدمة − المدفوعات، ويقدّر تعرض التأخر في التقديم والسداد.",
    "Opening GL VAT liability (credit = positive):": "التزام الضريبة الافتتاحي في الدفاتر (الدائن = موجب):", "Period End": "نهاية الفترة",
    "Net VAT (Box 13 + Box 14)": "صافي الضريبة (خانة 13 + خانة 14)", "Box 16 as Filed": "خانة 16 كما قُدمت", "Paid / (Refunded) for Period": "المسدد / (المسترد) للفترة",
    "Outstanding for Period": "المتبقي للفترة", "Penalty on Late Payments Made": "غرامة المدفوعات المتأخرة", "Penalty on Amount Still Unpaid (to As-of date)": "غرامة المبلغ غير المسدد (حتى تاريخ الاحتساب)",
    "GL VAT Liability at Period End": "التزام الضريبة في الدفاتر نهاية الفترة", "Expected Liability (Opening + Returns − Payments)": "الالتزام المتوقع (الافتتاحي + الإقرارات − المدفوعات)",
    "Difference GL − Expected": "الفرق: الدفاتر − المتوقع",
    "Penalties shown are indicative estimates only (5% of unpaid tax per month or part). ZATCA may also apply late-filing (5%–25%) and incorrect-return (50% of the difference) penalties – see Dashboard and Inspection_Guide.":
        f"الغرامات المعروضة تقديرية فقط (5% من الضريبة غير المسددة عن كل شهر أو جزء منه). قد تطبق الهيئة أيضاً غرامة التأخر في التقديم (5%–25%) وغرامة الإقرار الخاطئ (50% من الفرق) – راجع {S['Dashboard']} و{S['Inspection_Guide']}.",
    # Guide
    "ZATCA VAT Inspection – Practical Guide": "فحص ضريبة القيمة المضافة – دليل عملي",
    "Rules summarised from the KSA VAT Law & Implementing Regulations. Confirm current rates, penalties and deadlines with ZATCA or your advisor before relying on them.": "ملخص من نظام ضريبة القيمة المضافة ولائحته التنفيذية. تأكد من النسب والغرامات والمواعيد الحالية مع الهيئة أو مستشارك قبل الاعتماد عليها.",
    "HOW TO USE THIS WORKBOOK": "طريقة استخدام الملف",
    "Setup: enter company details, the audit year start date and filing frequency. Periods generate automatically.": f"{S['Setup']}: أدخل بيانات المنشأة وبداية سنة الفحص ودورية الإقرار. تُنشأ الفترات تلقائياً.",
    "TB_Monthly: paste opening balances and monthly net movements (Dr +, Cr −). Map EVERY account in column D – this drives all GL reconciliations.": f"{S['TB_Monthly']}: الصق الأرصدة الافتتاحية وصافي الحركات الشهرية (مدين +، دائن −). صنّف كل حساب في العمود D – فهو أساس جميع مطابقات الدفاتر.",
    "VAT_Returns: copy every filed return box-by-box from the ZATCA portal, including filing date. Check columns AS–AU.": f"{S['VAT_Returns']}: انقل كل إقرار مقدم خانةً خانة من بوابة الهيئة مع تاريخ التقديم. راجع الأعمدة AS–AU.",
    "Sales_Register / Purchase_Register: export invoice-level data from the ERP or FATOORA. For very high volumes (POS), use daily summary lines per VAT category.": f"{S['Sales_Register']} / {S['Purchase_Register']}: صدّر بيانات الفواتير من نظام ERP أو فاتورة. للأحجام الكبيرة (نقاط البيع) استخدم سطراً ملخصاً يومياً لكل فئة ضريبية.",
    "ZATCA_Payments: enter each SADAD payment and any refunds (negative).": f"{S['ZATCA_Payments']}: أدخل كل عملية سداد وأي مبالغ مستردة (بالسالب).",
    "Work through every ✗ on Recon_Sales, Recon_OutputVAT, Recon_InputVAT and Recon_VAT_Account. Document each explanation.": f"عالج كل علامة ✗ في {S['Recon_Sales']} و{S['Recon_OutputVAT']} و{S['Recon_InputVAT']} و{S['Recon_VAT_Account']}، ووثّق تفسير كل فرق.",
    "Use the Checklist to collect evidence; the Dashboard shows overall readiness and indicative exposure.": f"استخدم {S['Checklist']} لجمع الإثباتات؛ وتعرض {S['Dashboard']} الجاهزية العامة والتعرض التقديري.",
    "WHAT ZATCA TYPICALLY TESTS": "ما تفحصه الهيئة عادةً",
    "Revenue in audited financial statements / Zakat return vs total sales in VAT returns (the #1 test).": "الإيرادات في القوائم المالية المدققة / إقرار الزكاة مقابل إجمالي المبيعات في الإقرارات (الاختبار الأهم).",
    "Bank receipts vs declared sales; customs import data vs Box 8; withholding-tax returns vs Box 9 (reverse charge).": "المتحصلات البنكية مقابل المبيعات المُقر بها؛ بيانات الاستيراد الجمركية مقابل خانة 8؛ إقرارات ضريبة الاستقطاع مقابل خانة 9 (الاحتساب العكسي).",
    "Zero-rated and export sales: proof the goods left KSA (customs exit) or the service conditions were met.": "المبيعات بنسبة صفر والصادرات: إثبات خروج السلع من المملكة (بيان خروج جمركي) أو استيفاء شروط الخدمة.",
    "Input VAT: valid tax invoice in the company's name with a valid supplier VAT number; no blocked items.": "ضريبة المدخلات: فاتورة ضريبية صحيحة باسم المنشأة برقم ضريبي صحيح للمورد، دون بنود غير قابلة للخصم.",
    "Credit notes: valid reason, issued correctly, linked to the original invoice.": "الإشعارات الدائنة: سبب صحيح، وإصدار سليم، وربط بالفاتورة الأصلية.",
    "Deemed supplies: gifts, samples, own use, staff benefits, asset write-offs.": "التوريدات الافتراضية: الهدايا والعينات والاستخدام الخاص ومنافع الموظفين وشطب الأصول.",
    "E-invoicing: invoices generated/cleared/reported through a compliant solution with required fields and QR code.": "الفوترة الإلكترونية: إصدار الفواتير واعتمادها أو الإبلاغ عنها عبر حل متوافق يتضمن الحقول المطلوبة ورمز QR.",
    "Timeliness of filing and payment; use of Box 14 above the SAR 5,000 limit.": "الالتزام بمواعيد التقديم والسداد؛ واستخدام خانة 14 بما يتجاوز حد 5,000 ريال.",
    "COMMON FINDINGS IN KSA VAT AUDITS": "الملاحظات الشائعة في فحوصات ضريبة القيمة المضافة",
    "Revenue booked in GL but not declared (accruals, other income, recharges, asset sales).": "إيرادات مسجلة في الدفاتر وغير مُقر بها (إيرادات مستحقة، إيرادات أخرى، إعادة تحميل، بيع أصول).",
    "Exports zero-rated without exit evidence → reassessed at 15%.": "صادرات بنسبة صفر دون إثبات خروج ← يُعاد تقييمها بنسبة 15%.",
    "Input VAT claimed on invoices without supplier VAT number, in another entity's name, or on simplified invoices above limits.": "خصم ضريبة مدخلات على فواتير دون رقم ضريبي للمورد، أو باسم منشأة أخرى، أو على فواتير مبسطة تتجاوز الحدود.",
    "Reverse charge not applied on foreign services (consulting, software, royalties, management fees).": "عدم تطبيق الاحتساب العكسي على الخدمات الأجنبية (الاستشارات، البرمجيات، الإتاوات، أتعاب الإدارة).",
    "Customs VAT claimed where the company is not the importer of record.": "خصم ضريبة الجمارك رغم أن المنشأة ليست المستورد المسجل.",
    "Blocked input VAT claimed: entertainment, hospitality, private-use vehicles, staff personal benefits.": "خصم ضريبة غير قابلة للخصم: الترفيه والضيافة والسيارات للاستخدام الشخصي ومنافع الموظفين الشخصية.",
    "Errors above SAR 5,000 corrected through Box 14 instead of a voluntary disclosure.": "تصحيح أخطاء تتجاوز 5,000 ريال عبر خانة 14 بدلاً من الإفصاح الطوعي.",
    "KEY RULES & PENALTIES (verify current position)": "الأحكام والغرامات الرئيسية (تحقق من الوضع الحالي)",
    "Filing": "تقديم الإقرار", "Late filing": "التأخر في التقديم", "Late payment": "التأخر في السداد", "Incorrect return": "الإقرار الخاطئ", "Tax evasion": "التهرب الضريبي",
    "Corrections": "التصحيحات", "Records": "حفظ السجلات", "Assessment window": "مدة التقادم", "Objection": "الاعتراض", "Relief": "الإعفاء من الغرامات",
    "Monthly if taxable supplies exceed SAR 40 million a year, otherwise quarterly. Return and payment due by the last day of the following month.": "شهرياً إذا تجاوزت التوريدات الخاضعة 40 مليون ريال سنوياً، وإلا ربع سنوي. يستحق الإقرار والسداد في آخر يوم من الشهر التالي.",
    "5% to 25% of the tax due (VAT Law Art. 43).": "من 5% إلى 25% من الضريبة المستحقة (المادة 43).",
    "5% of the unpaid tax for each month or part of a month (VAT Law Art. 43).": "5% من الضريبة غير المسددة عن كل شهر أو جزء منه (المادة 43).",
    "50% of the difference between the tax calculated and the tax due (VAT Law Art. 42).": "50% من الفرق بين الضريبة المحتسبة والضريبة المستحقة (المادة 42).",
    "Up to three times the value of the goods or services concerned (VAT Law Art. 40).": "حتى ثلاثة أضعاف قيمة السلع أو الخدمات محل التهرب (المادة 40).",
    "Net errors above SAR 5,000 → voluntary disclosure (within 20 days of discovery). Errors up to SAR 5,000 can be corrected in Box 14 of the next return.": "الأخطاء التي يتجاوز صافيها 5,000 ريال ← إفصاح طوعي (خلال 20 يوماً من اكتشافها). الأخطاء حتى 5,000 ريال يمكن تصحيحها في خانة 14 من الإقرار التالي.",
    "Keep records at least 6 years; capital assets 11 years; real estate 15 years.": "احفظ السجلات 6 سنوات على الأقل؛ والأصول الرأسمالية 11 سنة؛ والعقارات 15 سنة.",
    "ZATCA can generally assess within 5 years from the end of the tax period (10 years where no return was filed or in cases of evasion).": "يحق للهيئة عادةً التقييم خلال 5 سنوات من نهاية الفترة الضريبية (10 سنوات في حال عدم تقديم الإقرار أو التهرب).",
    "File an objection against an assessment within the statutory deadline (currently 60 days from notification – verify), then escalate to the GSTC committees if needed.": "قدّم الاعتراض على التقييم خلال المهلة النظامية (حالياً 60 يوماً من الإبلاغ – تحقق منها)، ثم التصعيد إلى لجان الفصل في المخالفات والمنازعات الضريبية عند الحاجة.",
    "Check whether a ZATCA penalty-waiver / amnesty initiative is currently active – these have been extended several times.": "تحقق مما إذا كانت مبادرة الإعفاء من الغرامات سارية حالياً – فقد مُددت عدة مرات.",
    "RESPONSE TIPS": "نصائح للرد على الهيئة",
    "Answer only what is requested, in writing, within the deadline in the ZATCA notice; ask for an extension in writing if needed.": "أجب عما هو مطلوب فقط، كتابةً، وخلال المهلة المحددة في إشعار الهيئة؛ واطلب التمديد كتابةً عند الحاجة.",
    "Never submit a reconciliation with unexplained differences – fix or explain every ✗ first.": "لا تقدّم مطابقة بها فروق غير مفسرة – عالج أو فسّر كل علامة ✗ أولاً.",
    "Keep one index of everything submitted (Checklist col H) and the date it was sent.": f"احتفظ بفهرس واحد لكل ما تقدمه ({S['Checklist']} عمود H) وتاريخ إرساله.",
    "If you find an error before ZATCA does, assess whether a voluntary disclosure reduces the penalty.": "إذا اكتشفت خطأً قبل الهيئة، فادرس ما إذا كان الإفصاح الطوعي يخفض الغرامة.",
    "Get your tax advisor to review the pack before submission.": "اطلب من مستشارك الضريبي مراجعة الملف قبل التقديم.",
    # Lists sheet headers
    "Sales VAT Category": "الفئة الضريبية للمبيعات", "Purchase VAT Category": "الفئة الضريبية للمشتريات", "TB VAT Mapping": "التصنيف الضريبي للحسابات",
    "Expense Nature": "طبيعة المصروف", "Yes/No": "نعم/لا", "Checklist Status": "حالة البند",
    # comments
    "On the ZATCA return, VAT on reverse-charge imports (Box 9) is added to BOTH output VAT and input VAT, so it nets to zero for fully taxable businesses. Confirm against your filed return.":
        "في إقرار الهيئة تُضاف ضريبة الاستيراد بالاحتساب العكسي (خانة 9) إلى ضريبة المخرجات والمدخلات معاً، فيكون صافيها صفراً للمنشآت الخاضعة بالكامل. تحقق من ذلك مقابل الإقرار المقدم.",
    "VAT returns and payment are due by the last day of the month following the end of the tax period.": "يستحق الإقرار والسداد في آخر يوم من الشهر التالي لنهاية الفترة الضريبية.",
})
for m in range(1, 13):
    T[f"Movement M{m}"] = f"حركة الشهر {m}"

KEEP = {"dd-mmm-yyyy", "mmm yyyy", "yyyy"}  # number-format codes inside formulas
MONTHS_AR = '"يناير","فبراير","مارس","أبريل","مايو","يونيو","يوليو","أغسطس","سبتمبر","أكتوبر","نوفمبر","ديسمبر"'
missing = set()


def tr_text(v):
    if v in T:
        return T[v]
    if re.search(r"[A-Za-z]{3,}", v) and not re.fullmatch(r"[A-Z]{2,4}-[\w-]+|[0-9a-f-]+|SADAD-\d+", v):
        missing.add(v)
    return v


def tr_formula(f):
    for en, ar in SHEETS.items():
        f = re.sub(r"(?<![A-Za-z_'])" + en + r"!", f"'{ar}'!", f)

    def lit(m):
        s = m.group(1)
        if s in KEEP or not re.search(r"[A-Za-z]", s):
            return m.group(0)
        return '"' + tr_text(s).replace('"', '""') + '"'
    f = re.sub(r'"((?:[^"]|"")*)"', lit, f)
    return f


wb = load_workbook(SRC)
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            v = c.value
            if isinstance(v, str):
                c.value = tr_formula(v) if v.startswith("=") else tr_text(v)
            if c.number_format == "dd-mmm-yyyy":
                c.number_format = "dd/mm/yyyy"
            elif c.number_format == "mmm-yy":
                c.number_format = "mm/yyyy"
            if c.comment:
                c.comment = Comment(tr_text(c.comment.text), "VAT Pack")
    for dv in ws.data_validations.dataValidation:
        if dv.formula1:
            dv.formula1 = tr_formula(dv.formula1)
        if dv.error:
            dv.error = "اختر قيمة من القائمة"
    ws.sheet_view.rightToLeft = True

# Period labels and the dashboard header use month names; build them in Arabic.
setup = wb["Setup"]
for r in range(19, 31):
    setup[f"B{r}"] = (f'=IF(C{r}="","",IF($B$8="شهري",CHOOSE(MONTH(C{r}),{MONTHS_AR})&" "&YEAR(C{r}),'
                      f'"الربع "&A{r}&" "&YEAR(C{r})))')
dash = wb["Dashboard"]
dash["A2"] = dash["A2"].value.replace("dd-mmm-yyyy", "dd/mm/yyyy")

for ws in wb.worksheets:
    ws.title = SHEETS[ws.title]

wb.save(OUT)
if missing:
    print("UNTRANSLATED:")
    for m in sorted(missing):
        print("  ", repr(m))
print("saved", OUT)
